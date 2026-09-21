#!/usr/bin/env python3
"""
build_scorecard.py
------------------
Processes Google Search Console (GSC) exports into the two data files the pipeline
reads:
  1. data/scorecard.json  - performance scorecard per page (category, CTR, position, trend, priority)
  2. data/audit.json      - strategy per page (primary keyword, intent, money page) + blog slugs + money pages

Brand, domain and URL paths come from config.json. Nothing brand-specific is hard-coded.

INPUT CONTRACT (adjust the constants below to match your own exports):
  - PAGE_INVENTORY_XLSX : a workbook with sheets:
        "pages"        col A = slug, B = title, C = meta_title, D = meta_description
        "money_pages"  col A = name, B = url
        "blog"         col A = slug
        "meta_data"    col B = meta_title, C = meta_description, D = meta_keywords, Q(16) = h1_tag
  - GSC_DIR/AGGREGATE_XLSX : GSC "Pages" sheet, col A = url, B = clicks, C = impressions, D = ctr, E = position
  - GSC_DIR/<monthly files> : same "Pages" sheet, col A = url, B = clicks (used for trend)

Run from project root:  python build_scorecard.py
Requires: openpyxl  (pip install openpyxl)
"""

import json
import re
import statistics
from pathlib import Path

import openpyxl

from optimizer.config_loader import load_config

CONFIG = load_config()
BRAND = CONFIG["brand_name"]
BASE_URL = CONFIG["base_url"].rstrip("/") + "/"
INFO_PATH = CONFIG["info_path"]
COURSE_KEYWORD_MAP = CONFIG.get("course_keyword_map", {})

# ---------------------------------------------------------------------------
# Paths and input files  (EDIT THESE to match your exports)
# ---------------------------------------------------------------------------
BASE = Path(__file__).resolve().parent
GSC_DIR = BASE / "gsc"
AGGREGATE_XLSX = GSC_DIR / "16-month.xlsx"
PAGE_INVENTORY_XLSX = BASE / "page-inventory.xlsx"
OUT_DIR = BASE / "data"
OUT_DIR.mkdir(exist_ok=True)

# (label, filename) pairs, oldest to newest, used for the trend calculation
MONTHLY_FILES = [
    # ("Mar 2026", "monthly-mar.xlsx"),
    # ("Apr 2026", "monthly-apr.xlsx"),
]

# Sheet names inside PAGE_INVENTORY_XLSX
SHEET_PAGES = "pages"
SHEET_MONEY = "money_pages"
SHEET_BLOG = "blog"
SHEET_META = "meta_data"

# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------
def normalise_ctr(val):
    if val is None:
        return 0.0
    if isinstance(val, str):
        val = val.strip().rstrip("%")
        try:
            return float(val)
        except ValueError:
            return 0.0
    v = float(val)
    return round(v * 100, 2) if v < 1.0 else round(v, 2)


def safe_float(val, default=0.0):
    try:
        return float(val)
    except (ValueError, TypeError):
        return default


def safe_int(val, default=0):
    try:
        return int(float(val))
    except (ValueError, TypeError):
        return default


# ---------------------------------------------------------------------------
# 1. Load page inventory
# ---------------------------------------------------------------------------
print("[1/7] Loading page inventory ...")
wb = openpyxl.load_workbook(str(PAGE_INVENTORY_XLSX), read_only=True, data_only=True)

pages = []
for row in wb[SHEET_PAGES].iter_rows(min_row=2, values_only=True):
    if row[0] is None:
        continue
    pages.append({
        "slug": str(row[0]).strip(),
        "title": str(row[1]).strip() if row[1] else "",
        "meta_title": str(row[2]).strip() if row[2] else "",
        "meta_description": str(row[3]).strip() if row[3] else "",
    })
print(f"   Found {len(pages)} pages")

# 2. Money pages
print("[2/7] Loading money pages ...")
money_pages = []
if SHEET_MONEY in wb.sheetnames:
    for row in wb[SHEET_MONEY].iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        money_pages.append({
            "name": str(row[0]).strip(),
            "url": str(row[1]).strip() if row[1] else "",
        })
print(f"   Found {len(money_pages)} money pages")

# 3. Blog slugs (for cannibalization)
print("[3/7] Loading blog slugs ...")
blog_slugs = []
if SHEET_BLOG in wb.sheetnames:
    for row in wb[SHEET_BLOG].iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        blog_slugs.append(str(row[0]).strip())
print(f"   Found {len(blog_slugs)} blog slugs")

# 4. Meta data (join on meta_title)
print("[4/7] Loading meta data ...")
meta_by_title = {}
if SHEET_META in wb.sheetnames:
    for row in wb[SHEET_META].iter_rows(min_row=2, values_only=True):
        if row is None or len(row) < 2 or row[1] is None:
            continue
        title = str(row[1]).strip()
        kw = str(row[3]).strip() if (len(row) > 3 and row[3]) else ""
        if kw == "NULL":
            kw = ""
        h1 = str(row[16]).strip() if (len(row) > 16 and row[16]) else ""
        meta_by_title[title] = {"meta_keywords": kw, "h1_tag": h1}

for p in pages:
    mt = meta_by_title.get(p["meta_title"], {})
    p["meta_keywords"] = mt.get("meta_keywords", "")
    p["h1_tag"] = mt.get("h1_tag", "")
wb.close()

# 5. Aggregate GSC
print("[5/7] Loading aggregate GSC data ...")
gsc_by_url = {}
if AGGREGATE_XLSX.exists():
    wb_a = openpyxl.load_workbook(str(AGGREGATE_XLSX), read_only=True, data_only=True)
    for row in wb_a["Pages"].iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        gsc_by_url[str(row[0]).strip()] = {
            "clicks": safe_int(row[1]),
            "impressions": safe_int(row[2]),
            "ctr": normalise_ctr(row[3]),
            "position": round(safe_float(row[4]), 2),
        }
    wb_a.close()
else:
    print(f"   WARNING: {AGGREGATE_XLSX} not found; scorecard metrics will be zero")
print(f"   Loaded {len(gsc_by_url)} GSC page entries")

# 6. Monthly GSC (for trend)
print("[6/7] Loading monthly GSC data ...")
monthly_clicks = {}
month_labels = []
for label, filename in MONTHLY_FILES:
    fpath = GSC_DIR / filename
    if not fpath.exists():
        print(f"   WARNING: {filename} not found, skipping")
        continue
    month_labels.append(label)
    wb_m = openpyxl.load_workbook(str(fpath), read_only=True, data_only=True)
    for row in wb_m["Pages"].iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        url = str(row[0]).strip()
        monthly_clicks.setdefault(url, {})[label] = safe_int(row[1])
    wb_m.close()
print(f"   Loaded monthly data for {len(month_labels)} months")

# 7. Compute metrics
print("[7/7] Computing scorecard metrics ...")

def find_gsc_url(slug_tail):
    for c in (f"{BASE_URL}{INFO_PATH}/{slug_tail}", f"{BASE_URL}{INFO_PATH}/{slug_tail}/"):
        if c in gsc_by_url:
            return c
    for url in gsc_by_url:
        if url.rstrip("/").endswith(f"/{INFO_PATH}/{slug_tail}"):
            return url
    return None


def compute_trend(monthly_dict, labels):
    vals = [monthly_dict.get(m, 0) for m in labels]
    if len(vals) < 4:
        return 0.0
    half = len(vals) // 2
    sp, sr = sum(vals[:half]), sum(vals[half:])
    if sp == 0:
        return 100.0 if sr > 0 else 0.0
    return round(((sr - sp) / sp) * 100, 1)


def trend_label(pct):
    return "Rising" if pct > 15 else "Declining" if pct < -15 else "Stable"


def position_band(pos):
    if pos <= 0:
        return "No Data"
    if pos <= 3:
        return "Top 3"
    if pos <= 10:
        return "Page 1 (4-10)"
    if pos <= 20:
        return "Page 2"
    if pos <= 50:
        return "Page 3+"
    return "Deep"


def expected_ctr(pos):
    if pos <= 1:
        return 30.0
    if pos <= 2:
        return 15.0
    if pos <= 3:
        return 10.0
    if pos <= 10:
        return 5.0
    if pos <= 20:
        return 1.5
    return 0.5


def impression_band(imp):
    if imp >= 100000:
        return "Very High"
    if imp >= 10000:
        return "High"
    if imp >= 1000:
        return "Medium"
    if imp >= 100:
        return "Low"
    return "Very Low"


entries = []
matched = 0
for p in pages:
    slug_tail = p["slug"].replace(f"{INFO_PATH}/", "").strip("/")
    gsc_url = find_gsc_url(slug_tail)
    if gsc_url:
        matched += 1
        g = gsc_by_url[gsc_url]
        clicks, imp, ctr, pos = g["clicks"], g["impressions"], g["ctr"], g["position"]
    else:
        clicks = imp = 0
        ctr = pos = 0.0
    monthly_dict = monthly_clicks.get(gsc_url, {}) if gsc_url else {}
    tp = compute_trend(monthly_dict, month_labels)
    entries.append({
        "url": f"{BASE_URL}{p['slug']}",
        "title": p["title"],
        "slug": p["slug"],
        "clicks_16m": clicks,
        "impressions": imp,
        "ctr": ctr,
        "position": pos,
        "monthly": {m: monthly_dict.get(m, 0) for m in month_labels},
        "trend_pct": tp,
        "trend_label": trend_label(tp),
        "pos_band": position_band(pos),
        "ctr_gap": round(expected_ctr(pos) - ctr, 2) if pos > 0 else 0.0,
        "imp_band": impression_band(imp),
        "category": "",
        "priority_score": 0,
        "action": "",
    })
print(f"   Matched {matched}/{len(pages)} pages to GSC data")

# Category assignment (needs medians)
all_clicks = [e["clicks_16m"] for e in entries if e["clicks_16m"] > 0]
all_imp = [e["impressions"] for e in entries if e["impressions"] > 0]
median_clicks = statistics.median(all_clicks) if all_clicks else 0
p50_imp = statistics.median(all_imp) if all_imp else 0

for e in entries:
    clicks, imp, pos, tp = e["clicks_16m"], e["impressions"], e["position"], e["trend_pct"]
    if clicks < 10 and imp < 500:
        cat = "Dead Since Birth"
    elif clicks > median_clicks and tp < -15:
        cat = "Top Performers Falling"
    elif clicks > median_clicks and tp >= 0:
        cat = "Performing Well"
    elif imp > p50_imp and 8 <= pos <= 20:
        cat = "Quick Win"
    elif tp < -15:
        cat = "Lost Momentum"
    else:
        cat = "Performing Well" if clicks > 0 and tp >= 0 else "Lost Momentum" if clicks > 0 else "Dead Since Birth"
    e["category"] = cat
    e["action"] = {"Performing Well": "MAINTAIN", "Top Performers Falling": "PROTECT"}.get(cat, "OPTIMIZE")

# Priority score
max_imp = max((e["impressions"] for e in entries), default=1) or 1
max_gap = max((e["ctr_gap"] for e in entries), default=1) or 1
for e in entries:
    imp_score = min((e["impressions"] / max_imp) * 30, 30)
    gap_score = min((max(e["ctr_gap"], 0) / max_gap) * 25, 25) if max_gap > 0 else 0
    pos = e["position"]
    pos_score = 0 if pos <= 0 else 5 if pos <= 3 else 15 if pos <= 10 else 25 if pos <= 20 else 18 if pos <= 50 else 8
    tp = e["trend_pct"]
    trend_score = 20 if tp < -30 else 15 if tp < -15 else 10 if tp < 0 else 5 if tp < 15 else 2
    e["priority_score"] = min(round(imp_score + gap_score + pos_score + trend_score), 100)

entries.sort(key=lambda x: x["priority_score"], reverse=True)

cat_counts = {}
for e in entries:
    cat_counts[e["category"]] = cat_counts.get(e["category"], 0) + 1

with open(OUT_DIR / "scorecard.json", "w", encoding="utf-8") as f:
    json.dump({"pages": entries, "stats": {"total": len(pages), "in_gsc": matched, "categories": cat_counts}},
              f, indent=2, ensure_ascii=False)
print(f"   Saved {OUT_DIR / 'scorecard.json'}  | categories: {cat_counts}")

# ===========================================================================
# audit.json
# ===========================================================================
print("[audit] Building audit.json ...")

_brand_re = re.escape(BRAND)
BRAND_PATTERNS = re.compile(rf"\s*[\|\-–—]\s*{_brand_re}.*$|\b{_brand_re}\b", re.I)
SUFFIX_PATTERNS = re.compile(
    r"\s*[-–—]\s*(A Complete Guide|Complete Guide|Overview|"
    r"A Comprehensive Guide|Comprehensive Guide|An Ultimate Guide|Ultimate Guide|"
    r"A Beginner'?s Guide|Beginner'?s Guide|Everything You Need to Know|"
    r"All You Need to Know|A Detailed Guide|Detailed Guide|Tips & Tricks|"
    r"Key Insights|Top \d+|Best Practices).*$", re.I)
YEAR_PATTERN = re.compile(r"\s*\(?\s*20\d{2}\s*\)?\s*$")
LEADING_NUMBERS = re.compile(r"^\d+\+?\s*")
PIPE_SUFFIX = re.compile(r"\s*\|.*$")
COLON_SUFFIX = re.compile(r"\s*:\s*\d+.*$")


def derive_primary_keyword(meta_keywords, meta_title, title):
    if meta_keywords and meta_keywords.strip():
        parts = [p.strip() for p in meta_keywords.split(",") if p.strip()]
        if parts:
            return parts[0]
    kw = (meta_title or title or "").strip()
    if not kw:
        return ""
    kw = BRAND_PATTERNS.sub("", kw).strip()
    kw = PIPE_SUFFIX.sub("", kw).strip()
    kw = SUFFIX_PATTERNS.sub("", kw).strip()
    kw = COLON_SUFFIX.sub("", kw).strip()
    kw = YEAR_PATTERN.sub("", kw).strip()
    kw = LEADING_NUMBERS.sub("", kw).strip()
    return kw.strip(" -:|")


def match_money_page(slug, title):
    slug_l, title_l = slug.lower(), title.lower()
    best, best_score = None, 0
    stop = {"the", "a", "an", "in", "of", "for", "and", "to", "is", "are", "with",
            "on", "at", "by", "from", "or", "as", "training", "certification",
            "course", "exam", "questions", "answers", "salary", "interview",
            "practice", "test", "guide", "benefits", "vs", "roles", "what", "how",
            "why", "who", "when", "best", "top"}
    for mp in money_pages:
        name_l, url_l = mp["name"].lower(), mp["url"].lower()
        score = 0
        for key, terms in COURSE_KEYWORD_MAP.items():
            if key in slug_l:
                if any(t in name_l or t in url_l for t in terms):
                    score += 10
        overlap = (set(re.findall(r"[a-z]+", title_l)) - stop) & (set(re.findall(r"[a-z]+", name_l)) - stop)
        score += len(overlap) * 3
        if score > best_score:
            best, best_score = mp, score
    return best if best_score >= 3 else None


def classify_intent(slug, title):
    t = (slug + " " + title).lower()
    if any(s in t for s in ["vs", "versus", "comparison", "best", "top ", "tools", "software", "alternatives", "review"]):
        return "commercial"
    if any(s in t for s in ["login", "portal", "download", "register", "sign up"]):
        return "navigational"
    return "informational"


sc_lookup = {e["slug"]: e for e in entries}
audit_rows = []
for p in pages:
    sc = sc_lookup.get(p["slug"], {})
    mp = match_money_page(p["slug"], p["title"])
    audit_rows.append({
        "url": f"{BASE_URL}{p['slug']}",
        "slug": p["slug"],
        "title": p["title"],
        "meta_title": p["meta_title"],
        "meta_description": p["meta_description"],
        "h1_tag": p.get("h1_tag", ""),
        "meta_keywords": p.get("meta_keywords", ""),
        "primary_keyword": derive_primary_keyword(p.get("meta_keywords", ""), p.get("meta_title", ""), p.get("title", "")),
        "search_intent": classify_intent(p["slug"], p["title"]),
        "money_page": mp["url"] if mp else None,
        "money_page_name": mp["name"] if mp else None,
        "category": sc.get("category", ""),
        "priority_score": sc.get("priority_score", 0),
        "action": sc.get("action", ""),
    })
audit_rows.sort(key=lambda x: x["priority_score"], reverse=True)

with open(OUT_DIR / "audit.json", "w", encoding="utf-8") as f:
    json.dump({"audit_rows": audit_rows, "money_pages": money_pages, "blog_pages": blog_slugs},
              f, indent=2, ensure_ascii=False)
print(f"   Saved {OUT_DIR / 'audit.json'}  | rows: {len(audit_rows)}, money: {len(money_pages)}, blog: {len(blog_slugs)}")
print("Done.")
