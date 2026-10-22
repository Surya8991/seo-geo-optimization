"""Smoke test for build_scorecard.main(): importing it must not run the pipeline, and
a tiny synthetic inventory + GSC export must produce both data files."""
import json

import openpyxl
import pytest

import build_scorecard


def _inventory(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "pages"
    ws.append(["slug", "title", "meta_title", "meta_description"])
    ws.append(["info/what-is-agile", "What is Agile", "What is Agile | Your Brand", "Agile basics."])
    ws.append(["info/best-scrum-tools", "Best Scrum Tools", "Best Scrum Tools 2026", "Top tools."])
    money = wb.create_sheet("money_pages")
    money.append(["name", "url"])
    money.append(["Agile Course", "https://www.example.com/agile-course/"])
    blog = wb.create_sheet("blog")
    blog.append(["slug"])
    blog.append(["agile-tips"])
    wb.save(path)


def _aggregate(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pages"
    ws.append(["url", "clicks", "impressions", "ctr", "position"])
    ws.append(["https://www.example.com/info/what-is-agile", 120, 4000, 0.03, 9.5])
    wb.save(path)


def test_import_has_no_side_effects():
    assert callable(build_scorecard.main)


def test_main_builds_scorecard_and_audit(tmp_path):
    inv = tmp_path / "inventory.xlsx"
    gsc = tmp_path / "gsc"
    gsc.mkdir()
    _inventory(inv)
    _aggregate(gsc / "16-month.xlsx")
    out = tmp_path / "out"

    rc = build_scorecard.main(["--inventory", str(inv), "--gsc-dir", str(gsc), "--out-dir", str(out)])
    assert rc == 0

    scorecard = json.loads((out / "scorecard.json").read_text(encoding="utf-8"))
    assert scorecard["stats"]["total"] == 2
    assert scorecard["stats"]["in_gsc"] == 1
    agile = next(p for p in scorecard["pages"] if p["slug"] == "info/what-is-agile")
    assert agile["clicks_16m"] == 120 and agile["ctr"] == 3.0

    audit = json.loads((out / "audit.json").read_text(encoding="utf-8"))
    rows = {r["slug"]: r for r in audit["audit_rows"]}
    assert rows["info/what-is-agile"]["search_intent"] == "informational"
    assert rows["info/best-scrum-tools"]["search_intent"] == "commercial"
    assert rows["info/what-is-agile"]["money_page"] == "https://www.example.com/agile-course/"
    assert audit["blog_pages"] == ["agile-tips"]


def test_missing_inventory_returns_error(tmp_path, capsys):
    rc = build_scorecard.main(["--inventory", str(tmp_path / "nope.xlsx"), "--out-dir", str(tmp_path)])
    assert rc == 1
    assert "not found" in capsys.readouterr().err


@pytest.mark.parametrize("flag", ["--help"])
def test_help_exits_cleanly(flag):
    with pytest.raises(SystemExit) as e:
        build_scorecard.main([flag])
    assert e.value.code == 0
